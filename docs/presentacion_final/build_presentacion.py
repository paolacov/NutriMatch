"""Arma las 9 láminas de la compra. La narración va a las notas, no al lienzo."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
CAP = ROOT / "capturas"
ASSET = ROOT / "assets"

W, H = 13.333, 7.5
FONT = "Plus Jakarta Sans"

CREAM = "F7F3EB"
INK = "10231F"
TIDE = "0F6B63"
DEEP = "0A4A40"
BRICK = "C41E3A"
EMBER = "E86A2A"
PAPER = "FFFDF8"
SAND = "E6DFD3"
MUTE = "5E6864"
WHITE = "FFFFFF"
MIST = "E7F2F0"


NARRACION = [
    (
        "NutriMatch",
        "Cuando una persona compra un alimento empacado, tiene que revisar piezas distintas "
        "para comparar opciones. NutriMatch organiza esa información para que la decisión sea más clara.",
    ),
    (
        "Elegir un empacado reparte la información",
        "La información necesaria para decidir puede estar repartida entre nutrición, procesamiento, "
        "etiquetas, restricciones y precio. Si una pieza falta, NutriMatch no la interpreta como un valor "
        "favorable ni como uno desfavorable. Esa es la condición para comparar de forma consistente.",
    ),
    (
        "La persona dice qué le importa",
        "Antes de comparar, la persona declara sus prioridades. En esta demostración el orden es "
        "nutrición, procesamiento y etiquetas. No hay dieta ni alergias. El resultado depende de ese "
        "orden. No significa que el sistema determine qué alimento es saludable.",
    ),
    (
        "Encuentra el producto que tienes enfrente",
        "Para esta demostración parto de un producto concreto. Escribo el código 7501003390288 en Buscar. "
        "Ese código abre directamente la ficha de aceitunas sin hueso, de Búfalo.",
    ),
    (
        "Una sola vista para entender el producto",
        "La ficha reúne la información relevante para la decisión. Con las prioridades declaradas, "
        "Búfalo queda cerca de 40. Los 30 pesos son un precio real. En la foto se ven los sellos de exceso "
        "de grasas saturadas y de sodio. El resultado no significa que el alimento sea saludable.",
    ),
    (
        "Otras opciones del mismo grupo",
        "La persona no tiene que quedarse con el primer producto. Al explorar alternativas del mismo grupo "
        "aparece La Cibeles. Con las prioridades declaradas, obtiene un resultado mayor: la pantalla muestra "
        "cerca de 76,5. Eso no significa que sea más saludable. Los 74 pesos son un precio de demostración, "
        "porque no hay un precio real para ese producto. El precio se muestra como información y no determina "
        "el orden. La acción de esta historia es Sustituir.",
    ),
    (
        "La elección queda en el carrito",
        "Después de pulsar Sustituir, Búfalo deja el carrito y queda únicamente La Cibeles, con la etiqueta "
        "de precio de demostración. La Cibeles no pertenece a los tres grupos de esta vista del Plato del "
        "Buen Comer, por eso el carrito aparece como no clasificado. Con un solo producto, ese grupo queda "
        "en 100 por ciento. El porcentaje se calcula por número de productos, no por gramos.",
    ),
    (
        "Las alertas informan",
        "Después de la sustitución, las alertas corresponden al carrito actual. La Cibeles aparece sin sello "
        "de exceso. Las alertas informan: no constituyen un diagnóstico y no determinan por sí mismas la puntuación.",
    ),
    (
        "Detrás de esa comparación",
        "Detrás de la experiencia hay un flujo de Ciencia de Datos: integra los datos, los prepara, construye "
        "las variables de la decisión y alimenta un motor determinista. 5 864 productos están dentro del "
        "universo puntuable: cumplen las condiciones del motor para poder entrar al ranking. Un dato faltante "
        "no se convierte en cero, y no se imputaron nutrientes, alérgenos ni precio. "
        "NutriMatch utiliza un motor determinista porque el catálogo no contiene un target observado que "
        "permita entrenar y evaluar de manera científicamente válida un modelo de recomendación.\n\n"
        "El valor de NutriMatch está en transformar información dispersa en una comparación trazable y "
        "utilizable. No decide por la persona qué alimento es saludable. Organiza la información para que "
        "pueda comparar y elegir.",
    ),
]


def rgb(hex_color):
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def frame(im, radius=28):
    im = im.convert("RGBA")
    mask = Image.new("L", im.size, 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, im.width - 1, im.height - 1), radius=radius, fill=255)
    im.putalpha(mask)
    pad = 24
    canvas = Image.new("RGBA", (im.width + pad * 2, im.height + pad * 2), (0, 0, 0, 0))
    shadow = Image.new("RGBA", im.size, (16, 35, 31, 255))
    shadow.putalpha(mask.point(lambda a: 60 if a > 0 else 0))
    shadow = shadow.filter(ImageFilter.GaussianBlur(10))
    canvas.paste(shadow, (pad + 2, pad + 8), shadow)
    canvas.paste(im, (pad, pad), im)
    # RGB sin canal alfa: Keynote rechaza varios PNG con transparencia.
    cream = Image.new("RGB", canvas.size, (247, 243, 235))
    cream.paste(canvas, mask=canvas.split()[-1])
    return cream


def crop(src, box, dest):
    im = Image.open(src).crop(box)
    if im.width > 2200:
        nh = int(im.height * 2200 / im.width)
        im = im.resize((2200, nh), Image.Resampling.LANCZOS)
    frame(im).save(dest)
    return dest


def prepare_images():
    ASSET.mkdir(exist_ok=True)
    shots = {
        "home": crop(CAP / "buscar.jpg", (155, 155, 760, 1070), ASSET / "s1-nuti.png"),
        "hero": crop(CAP / "ficha-bufalo.jpg", (210, 400, 1670, 615), ASSET / "s1-hero.png"),
        "buscar": crop(CAP / "buscar-codigo.png", (250, 806, 2680, 985), ASSET / "s4-buscar.png"),
        "ficha": crop(CAP / "ficha-bufalo.jpg", (140, 150, 1760, 1285), ASSET / "s5-ficha.png"),
        "alt": crop(CAP / "alternativas.jpg", (248, 575, 1008, 1015), ASSET / "s6-alt.png"),
        "cart": crop(CAP / "carrito-cibeles.png", (240, 210, 2460, 650), ASSET / "s7-cart.png"),
        "plato_l": crop(CAP / "plato-100.png", (330, 355, 1140, 720), ASSET / "s7-plato-l.png"),
        "plato_r": crop(CAP / "plato-100.png", (1180, 410, 2400, 750), ASSET / "s7-plato-r.png"),
        "alertas": crop(CAP / "alertas-cibeles.png", (180, 210, 2050, 1740), ASSET / "s8-alertas.png"),
    }
    return shots


def set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rpr.find(qn(tag))
        if node is None:
            node = rpr.makeelement(qn(tag), {})
            rpr.append(node)
        node.set("typeface", FONT)


def add_text(slide, text, x, y, w, h, size, color, bold=False, align="left", anchor="top"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 0.95
        set_run(p.add_run(), line, size, color, bold)
    return shape


def rect(slide, x, y, w, h, fill, line=None, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1.25)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = radius
    return shape


def picture(slide, path, x, y, max_w, max_h):
    im = Image.open(path)
    aspect = im.height / im.width
    dw, dh = max_w, max_w * aspect
    if dh > max_h:
        dh = max_h
        dw = dh / aspect
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(dw), Inches(dh))
    return dw, dh


def background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(CREAM)


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide)
    return slide


def kicker(slide, number, section):
    add_text(slide, f"{number:02d}   /   {section}", 0.48, 0.28, 8, 0.28, 13, TIDE, bold=True)


def slide_number(slide, number):
    add_text(slide, f"{number:02d}", 12.15, 7.12, 0.7, 0.26, 12, TIDE, bold=True, align="right")


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text



def slide_portada(prs, shots, narration):
    s = new_slide(prs)
    rect(s, 0, 0, 0.18, H, TIDE)
    rect(s, 7.35, 0, 5.99, H, PAPER)
    rect(s, 7.35, 0, 0.08, H, EMBER)
    add_text(s, "DIPLOMADO EN CIENCIA DE DATOS  ·  2026", 0.5, 0.42, 6.5, 0.28, 12, TIDE, bold=True)
    add_text(s, "NutriMatch", 0.48, 1.15, 6.6, 0.95, 54, INK, bold=True)
    add_text(s, "Una forma más clara de decidir\nfrente al anaquel", 0.5, 2.25, 6.4, 1.15, 24, DEEP, bold=True)
    rect(s, 0.5, 3.6, 1.4, 0.07, EMBER)
    add_text(s, "Paola Castañeda", 0.5, 3.9, 6.2, 0.38, 18, INK, bold=True)
    steps = [
        ("01", "Encuentra", "el producto que tienes enfrente"),
        ("02", "Compara", "con las prioridades declaradas"),
        ("03", "Elige", "y revisa lo que informa el carrito"),
    ]
    for i, (num, title, line) in enumerate(steps):
        y = 4.55 + i * 0.8
        rect(s, 0.5, y, 6.5, 0.7, PAPER, SAND, radius=0.12)
        add_text(s, num, 0.68, y, 0.7, 0.7, 18, EMBER, bold=True, anchor="middle")
        add_text(s, title, 1.45, y + 0.06, 5.2, 0.32, 16, INK, bold=True)
        add_text(s, line, 1.45, y + 0.34, 5.2, 0.28, 13, MUTE)
    picture(s, shots["home"], 8.55, 0.18, 4.3, 5.2)
    picture(s, shots["hero"], 7.58, 5.55, 5.5, 1.55)
    notes(s, narration)


def slide_problema(prs, narration):
    s = new_slide(prs)
    kicker(s, 2, "PROBLEMA")
    add_text(s, "Elegir un empacado reparte la información", 0.42, 0.58, 12.4, 0.5, 28, INK, bold=True)
    add_text(s, "DECISIÓN DE COMPRA", 0.42, 1.2, 4.2, 0.28, 12, EMBER, bold=True)
    rect(s, 2.55, 1.32, 10.3, 0.015, SAND)
    cards = [
        ("Nutrición", "¿Qué aporta?", TIDE),
        ("Procesamiento", "¿Cómo fue procesado?", DEEP),
        ("Etiquetas", "¿Qué sellos tiene?", EMBER),
        ("Restricciones", "¿Hay algo que deba evitar?", BRICK),
        ("Precio", "¿Cuánto cuesta?", TIDE),
    ]
    for i, (title, question, color) in enumerate(cards):
        x = 0.42 + i * 2.56
        rect(s, x, 1.62, 2.4, 3.62, PAPER, SAND, radius=0.1)
        rect(s, x, 1.62, 2.4, 0.1, color)
        add_text(s, f"0{i + 1}", x + 0.16, 1.9, 2.05, 0.55, 26, color, bold=True)
        add_text(s, title, x + 0.16, 2.55, 2.08, 0.7, 16, INK, bold=True)
        add_text(s, question, x + 0.16, 3.4, 2.08, 1.05, 15, MUTE)
    rect(s, 0, 5.45, W, 2.05, DEEP)
    add_text(s, "Un dato ausente no es un cero.", 0.5, 5.85, 12.2, 0.7, 32, WHITE, bold=True)
    add_text(s, "Si falta un dato, no se lee como favorable ni como desfavorable.", 0.5, 6.6, 11, 0.35, 15, "D5E8E4")
    add_text(s, "02", 12.15, 6.95, 0.7, 0.28, 12, "9FCFC6", bold=True, align="right")
    notes(s, narration)


def slide_prioridades(prs, narration):
    s = new_slide(prs)
    kicker(s, 3, "PERSONA")
    slide_number(s, 3)
    add_text(s, "La persona dice qué le importa", 0.42, 0.58, 12, 0.5, 30, INK, bold=True)
    rows = [
        ("01", "Nutrición", TIDE, WHITE, 1.55),
        ("02", "Procesamiento", MIST, INK, 1.35),
        ("03", "Etiquetas", PAPER, INK, 1.2),
    ]
    y = 1.28
    for num, name, fill, color, height in rows:
        rect(s, 0.42, y, 7.7, height, fill, None if fill != PAPER else SAND, radius=0.1)
        add_text(s, num, 0.65, y, 1.3, height, 28, EMBER if fill != TIDE else WHITE, bold=True, anchor="middle")
        add_text(s, name, 2.15, y, 5.5, height, 28, color, bold=True, anchor="middle")
        y += height + 0.16
    rect(s, 8.4, 1.28, 4.5, 2.1, PAPER, SAND, radius=0.1)
    add_text(s, "DIETA", 8.65, 1.48, 4.05, 0.28, 12, TIDE, bold=True)
    add_text(s, "Ahora no", 8.65, 2.05, 4.05, 0.7, 28, INK, bold=True)
    rect(s, 8.4, 3.6, 4.5, 2.1, PAPER, SAND, radius=0.1)
    add_text(s, "ALERGIAS", 8.65, 3.8, 4.05, 0.28, 12, TIDE, bold=True)
    add_text(s, "Ahora no", 8.65, 4.37, 4.05, 0.7, 28, INK, bold=True)
    add_text(s, "La comparación parte de\nlas prioridades declaradas.", 0.5, 6.45, 8, 0.7, 16, DEEP, bold=True)
    notes(s, narration)


def slide_buscar(prs, shots, narration):
    s = new_slide(prs)
    kicker(s, 4, "BUSCAR")
    slide_number(s, 4)
    add_text(s, "Encuentra el producto que tienes enfrente", 0.42, 0.55, 12.4, 0.48, 28, INK, bold=True)
    cards = [
        ("PRODUCTO", "Aceitunas\nsin hueso", PAPER, INK, SAND),
        ("MARCA", "Búfalo", MIST, DEEP, None),
        ("CÓDIGO", "7501003390288", TIDE, WHITE, None),
    ]
    for i, (label, value, fill, color, line) in enumerate(cards):
        x = 0.42 + i * 4.28
        rect(s, x, 1.18, 4.08, 3.85, fill, line, radius=0.1)
        add_text(s, label, x + 0.24, 1.4, 3.6, 0.28, 12, EMBER if fill != TIDE else "D5E8E4", bold=True)
        add_text(s, value, x + 0.24, 2.15, 3.6, 1.7, 30, color, bold=True)
        rect(s, x + 0.24, 4.45, 1.15, 0.055, EMBER if fill != TIDE else WHITE)
    add_text(s, "Ese código abre la ficha.", 0.48, 5.18, 8, 0.32, 16, DEEP, bold=True)
    picture(s, shots["buscar"], 0.42, 5.55, 12.5, 1.6)
    notes(s, narration)


def slide_ficha(prs, shots, narration):
    s = new_slide(prs)
    kicker(s, 5, "FICHA")
    slide_number(s, 5)
    add_text(s, "Una sola vista para entender el producto", 0.42, 0.55, 12.4, 0.45, 26, INK, bold=True)
    dw, dh = picture(s, shots["ficha"], 2.42, 1.12, 7.7, 5.95)
    left = [
        ("≈ 40", "Resultado"),
        ("Nutrición", "37.1"),
        ("Procesamiento", "44.4"),
        ("Etiquetas", "Sin preferencias"),
    ]
    for i, (title, line) in enumerate(left):
        y = 1.28 + i * 1.28
        add_text(s, title, 0.28, y, 2.05, 0.4, 16 if i else 28, DEEP if i == 0 else INK, bold=True)
        add_text(s, line, 0.28, y + 0.38, 2.05, 0.28, 12, MUTE)
        rect(s, 2.15, y + 0.22, 0.22, 0.035, EMBER if i == 0 else SAND)
    rx = 2.42 + dw + 0.16
    add_text(s, "$30", rx, 2.15, 2.5, 0.55, 32, INK, bold=True)
    rect(s, rx, 2.8, 1.85, 0.36, TIDE, radius=0.16)
    add_text(s, "Precio real", rx, 2.84, 1.85, 0.28, 12, WHITE, bold=True, align="center")
    add_text(s, "Sellos de exceso", rx, 3.4, 2.5, 0.55, 15, INK, bold=True)
    add_text(s, "Grasas saturadas\nSodio", rx, 4.0, 2.5, 0.65, 14, MUTE)
    notes(s, narration)


def slide_alternativas(prs, shots, narration):
    s = new_slide(prs)
    kicker(s, 6, "ALTERNATIVAS")
    slide_number(s, 6)
    add_text(s, "Otras opciones del mismo grupo", 0.4, 0.52, 12, 0.42, 28, INK, bold=True)
    rect(s, 0.4, 1.15, 3.55, 4.55, PAPER, SAND, radius=0.1)
    add_text(s, "BÚFALO", 0.58, 1.32, 3.2, 0.28, 12, MUTE, bold=True)
    add_text(s, "≈ 40", 0.58, 1.65, 3.2, 0.6, 36, DEEP, bold=True)
    add_text(s, "$30", 0.58, 2.3, 3.2, 0.4, 22, INK, bold=True)
    rect(s, 0.58, 2.8, 1.7, 0.34, TIDE, radius=0.16)
    add_text(s, "Precio real", 0.58, 2.82, 1.7, 0.3, 12, WHITE, bold=True, align="center")
    add_text(s, "↓", 0.58, 3.25, 1, 0.32, 18, EMBER, bold=True)
    add_text(s, "Explorar alternativas", 0.58, 3.58, 3.2, 0.35, 14, INK, bold=True)
    add_text(s, "↓", 0.58, 3.95, 1, 0.32, 18, EMBER, bold=True)
    rect(s, 0.58, 4.4, 2.15, 0.4, TIDE, radius=0.16)
    add_text(s, "Sustituir", 0.58, 4.44, 2.15, 0.32, 14, WHITE, bold=True, align="center")
    picture(s, shots["alt"], 4.15, 1.1, 8.75, 4.65)
    add_text(s, "Con las prioridades que declaré, queda por encima.", 0.42, 5.95, 12.3, 0.4, 18, BRICK, bold=True)
    add_text(s, "El precio se muestra como información; no determina el orden.", 0.42, 6.45, 12, 0.32, 14, MUTE)
    notes(s, narration)


def slide_carrito(prs, shots, narration):
    s = new_slide(prs)
    kicker(s, 7, "CARRITO")
    slide_number(s, 7)
    add_text(s, "La elección queda en el carrito", 0.42, 0.52, 12, 0.42, 28, INK, bold=True)
    add_text(s, "ELECCIÓN", 0.42, 1.02, 2.2, 0.24, 12, TIDE, bold=True)
    add_text(s, "La Cibeles  ·  $74  ·  Precio de demostración", 2.5, 1.0, 8.5, 0.28, 14, EMBER, bold=True)
    _, cart_h = picture(s, shots["cart"], 0.38, 1.32, 12.55, 2.25)
    plate_y = 1.32 + cart_h + 0.32
    add_text(s, "DISTRIBUCIÓN", 0.42, plate_y, 2.6, 0.24, 12, TIDE, bold=True)
    add_text(s, "No clasificado  ·  100 %", 3.15, plate_y - 0.02, 7, 0.28, 16, DEEP, bold=True)
    picture(s, shots["plato_l"], 0.38, plate_y + 0.32, 4.55, 2.15)
    picture(s, shots["plato_r"], 5.05, plate_y + 0.32, 7.9, 2.15)
    add_text(s, "Distribución por número de productos", 0.42, 6.95, 8, 0.28, 13, MUTE)
    notes(s, narration)


def slide_alertas(prs, shots, narration):
    s = new_slide(prs)
    kicker(s, 8, "ALERTAS")
    slide_number(s, 8)
    add_text(s, "Las alertas informan", 0.42, 0.52, 8, 0.42, 28, INK, bold=True)
    picture(s, shots["alertas"], 0.35, 1.1, 8.15, 6.05)
    rect(s, 8.7, 1.35, 4.2, 5.4, PAPER, SAND, radius=0.1)
    add_text(s, "Sin sello\nde exceso", 8.95, 1.6, 3.75, 1.15, 26, DEEP, bold=True)
    add_text(s, "No traen un sello\nNOM-051 de exceso.", 8.95, 2.9, 3.75, 0.75, 15, INK, bold=True)
    rect(s, 8.95, 3.8, 1.3, 0.05, EMBER)
    add_text(s, "Informar ≠ diagnosticar", 8.95, 4.1, 3.75, 0.7, 18, BRICK, bold=True)
    add_text(s, "Las alertas no determinan\npor sí mismas la puntuación.", 8.95, 5.0, 3.75, 0.75, 14, MUTE)
    notes(s, narration)


def slide_detras(prs, narration):
    s = new_slide(prs)
    kicker(s, 9, "FUNDAMENTO")
    add_text(s, "Detrás de esa comparación", 0.42, 0.52, 12, 0.42, 28, INK, bold=True)
    rect(s, 0.42, 1.1, 6.15, 1.45, PAPER, SAND, radius=0.1)
    add_text(s, "13 093", 0.62, 1.2, 5.7, 0.7, 36, DEEP, bold=True)
    add_text(s, "productos integrados", 0.64, 1.95, 5.6, 0.35, 14, MUTE)
    rect(s, 6.75, 1.1, 6.15, 1.45, MIST, radius=0.1)
    add_text(s, "5 864", 6.95, 1.2, 5.7, 0.7, 36, TIDE, bold=True)
    add_text(s, "productos dentro del universo puntuable", 6.97, 1.95, 5.7, 0.35, 14, DEEP)
    steps = ["Datos", "Preparación", "Variables de decisión", "Motor determinista", "Alternativas"]
    rect(s, 0.6, 2.9, 0.045, 1.85, "C9DDD8")
    for i, name in enumerate(steps):
        y = 2.75 + i * 0.48
        rect(s, 0.42, y, 0.42, 0.38, TIDE, radius=0.15)
        add_text(s, str(i + 1), 0.42, y, 0.42, 0.38, 12, WHITE, bold=True, align="center", anchor="middle")
        add_text(s, name, 1.0, y, 4.5, 0.38, 15, INK, bold=True, anchor="middle")
    rect(s, 6.75, 2.75, 6.15, 2.35, PAPER, SAND, radius=0.1)
    add_text(s, "Un dato faltante no se convierte\nautomáticamente en cero.", 6.98, 3.05, 5.75, 1.0, 20, DEEP, bold=True)
    add_text(s, "El motor no lo trata como un valor\nfavorable ni desfavorable.", 6.98, 4.2, 5.75, 0.65, 14, MUTE)
    rect(s, 0, 5.35, W, 2.15, DEEP)
    add_text(
        s,
        "NutriMatch no decide por la persona\nqué alimento es saludable.",
        0.5, 5.5, 12.2, 1.05, 26, WHITE, bold=True,
    )
    add_text(
        s,
        "Organiza la información para que pueda comparar y elegir.",
        0.5, 6.65, 11.2, 0.35, 15, "D5E8E4",
    )
    add_text(s, "09", 12.15, 6.95, 0.7, 0.28, 12, "9FCFC6", bold=True, align="right")
    notes(s, narration)

def write_narracion():
    lines = [
        "# Narración de la presentación",
        "",
        "Estas notas no aparecen en las láminas. Están también en las notas del presentador del PPTX.",
        "Duración orientativa del recorrido: unos 5 minutos.",
        "",
    ]
    for i, (title, text) in enumerate(NARRACION, 1):
        lines.append(f"## {i}. {title}")
        lines.append("")
        lines.append(text)
        lines.append("")
    path = ROOT / "narracion_presentacion.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def build(shots):
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "NutriMatch"
    prs.core_properties.author = "Paola Castañeda"
    prs.core_properties.subject = "Decidir una compra con NutriMatch"
    builders = [
        lambda: slide_portada(prs, shots, NARRACION[0][1]),
        lambda: slide_problema(prs, NARRACION[1][1]),
        lambda: slide_prioridades(prs, NARRACION[2][1]),
        lambda: slide_buscar(prs, shots, NARRACION[3][1]),
        lambda: slide_ficha(prs, shots, NARRACION[4][1]),
        lambda: slide_alternativas(prs, shots, NARRACION[5][1]),
        lambda: slide_carrito(prs, shots, NARRACION[6][1]),
        lambda: slide_alertas(prs, shots, NARRACION[7][1]),
        lambda: slide_detras(prs, NARRACION[8][1]),
    ]
    for build_slide in builders:
        build_slide()
    repair_for_keynote(prs)
    out = ROOT / "presentacion_nutrimatch.pptx"
    prs.save(out)
    return out


def repair_for_keynote(prs):
    """Keynote no abre el archivo si el grupo raíz de la lámina no trae transformación."""
    from lxml import etree

    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    for slide in prs.slides:
        grp = slide.shapes._spTree.find(qn("p:grpSpPr"))
        if grp is None or grp.find(qn("a:xfrm")) is not None:
            continue
        xfrm = etree.SubElement(grp, f"{ns}xfrm")
        etree.SubElement(xfrm, f"{ns}off", x="0", y="0")
        etree.SubElement(xfrm, f"{ns}ext", cx="0", cy="0")
        etree.SubElement(xfrm, f"{ns}chOff", x="0", y="0")
        etree.SubElement(xfrm, f"{ns}chExt", cx="0", cy="0")


def export_pdf(pptx_path):
    """Dibuja el PPTX a PDF con Plus Jakarta Sans incrustada. No incluye las notas."""
    from io import BytesIO

    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from reportlab.lib.colors import Color
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.platypus import Paragraph

    fonts = ROOT / "fuentes"
    pdfmetrics.registerFont(TTFont("PJ", str(fonts / "PlusJakartaSans-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("PJ-Bold", str(fonts / "PlusJakartaSans-Bold.ttf")))

    prs = Presentation(str(pptx_path))
    page_w = prs.slide_width / 914400 * 72
    page_h = prs.slide_height / 914400 * 72
    out = ROOT / "presentacion_nutrimatch.pdf"
    c = pdfcanvas.Canvas(str(out), pagesize=(page_w, page_h))
    c.setTitle("NutriMatch")
    c.setAuthor("Paola Castañeda")

    def pt(emu):
        return emu / 914400 * 72

    def color_of(rgb_color):
        return Color(rgb_color[0] / 255, rgb_color[1] / 255, rgb_color[2] / 255)

    def esc(text):
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    for slide in prs.slides:
        c.setFillColor(color_of(rgb(CREAM)))
        c.rect(0, 0, page_w, page_h, fill=1, stroke=0)
        for shape in slide.shapes:
            x, y = pt(shape.left), pt(shape.top)
            w, h = pt(shape.width), pt(shape.height)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                c.drawImage(
                    ImageReader(BytesIO(shape.image.blob)),
                    x,
                    page_h - y - h,
                    w,
                    h,
                    mask="auto",
                    preserveAspectRatio=False,
                )
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                fill = None
                try:
                    fill = color_of(shape.fill.fore_color.rgb)
                except Exception:
                    fill = None
                stroke = None
                try:
                    if shape.line.fill.type is not None:
                        stroke = color_of(shape.line.color.rgb)
                except Exception:
                    stroke = None
                radius = 0
                if shape.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
                    try:
                        radius = min(w, h) * float(shape.adjustments[0])
                    except Exception:
                        radius = min(w, h) * 0.15
                c.saveState()
                if fill:
                    c.setFillColor(fill)
                if stroke:
                    c.setStrokeColor(stroke)
                    c.setLineWidth(1.25)
                if radius > 0:
                    c.roundRect(x, page_h - y - h, w, h, radius, fill=1 if fill else 0, stroke=1 if stroke else 0)
                else:
                    c.rect(x, page_h - y - h, w, h, fill=1 if fill else 0, stroke=1 if stroke else 0)
                c.restoreState()
            if not getattr(shape, "has_text_frame", False):
                continue
            paragraphs = []
            font_name = "PJ"
            size = 14
            text_color = color_of(rgb(INK))
            align = TA_LEFT
            for p in shape.text_frame.paragraphs:
                text = "".join(run.text for run in p.runs)
                if p.runs:
                    run = p.runs[0]
                    if run.font.size is not None:
                        size = run.font.size.pt
                    font_name = "PJ-Bold" if run.font.bold else "PJ"
                    try:
                        text_color = color_of(run.font.color.rgb)
                    except Exception:
                        pass
                    if p.alignment == PP_ALIGN.CENTER:
                        align = TA_CENTER
                    elif p.alignment == PP_ALIGN.RIGHT:
                        align = TA_RIGHT
                    else:
                        align = TA_LEFT
                paragraphs.append(esc(text))
            if not any(paragraphs):
                continue
            style = ParagraphStyle(
                "s",
                fontName=font_name,
                fontSize=size,
                leading=size * 1.2,
                textColor=text_color,
                alignment=align,
            )
            para = Paragraph("<br/>".join(paragraphs), style)
            _pw, ph = para.wrap(w, h + 40)
            anchor = shape.text_frame.vertical_anchor
            if anchor == MSO_ANCHOR.MIDDLE:
                yb = page_h - y - h + max(0, (h - ph) / 2)
            elif anchor == MSO_ANCHOR.BOTTOM:
                yb = page_h - y - h
            else:
                yb = page_h - y - ph
            para.drawOn(c, x, yb)
        c.showPage()
    c.save()
    return out


def rewrite_notes_for_keynote(path):
    """Sustituye las notas de python-pptx, que Keynote rechaza, por el formato que sí abre."""
    import zipfile

    from lxml import etree

    prs = Presentation(str(path))
    notes = [slide.notes_slide.notes_text_frame.text for slide in prs.slides]
    for slide in prs.slides:
        for rel in list(slide.part.rels.values()):
            if "notesSlide" in rel.reltype:
                slide.part.drop_rel(rel.rId)
    for rel in list(prs.part.rels.values()):
        if "notesMaster" in rel.reltype:
            prs.part.drop_rel(rel.rId)
    prs.save(str(path))

    template = ROOT / "notas_keynote"
    master_xml = (template / "notesMaster1.xml").read_bytes()
    theme2 = (template / "theme2.xml").read_bytes()

    def esc(text):
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def notes_xml(text):
        paragraphs = []
        for para in text.split("\n"):
            if para.strip():
                paragraphs.append(f"<a:p><a:pPr/><a:r><a:t>{esc(para)}</a:t></a:r></a:p>")
            else:
                paragraphs.append("<a:p><a:pPr/></a:p>")
        body = "".join(paragraphs)
        return f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:notes xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
<p:cSld><p:spTree>
<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>
<p:sp><p:nvSpPr><p:cNvPr id="2" name="Slide Image"/><p:cNvSpPr/><p:nvPr><p:ph type="sldImg"/></p:nvPr></p:nvSpPr>
<p:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:pPr/></a:p></p:txBody></p:sp>
<p:sp><p:nvSpPr><p:cNvPr id="3" name="Notes"/><p:cNvSpPr/><p:nvPr><p:ph type="body" sz="quarter" idx="1"/></p:nvPr></p:nvSpPr>
<p:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
<p:txBody><a:bodyPr/><a:lstStyle/>{body}</p:txBody></p:sp>
</p:spTree></p:cSld>
<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:notes>
""".encode()

    nsr = "http://schemas.openxmlformats.org/package/2006/relationships"
    nsct = "http://schemas.openxmlformats.org/package/2006/content-types"
    nsp = "http://schemas.openxmlformats.org/presentationml/2006/main"
    nsr_attr = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    with zipfile.ZipFile(path) as zin:
        files = {name: zin.read(name) for name in zin.namelist()}

    files["ppt/theme/theme2.xml"] = theme2
    files["ppt/notesMasters/notesMaster1.xml"] = master_xml
    files["ppt/notesMasters/_rels/notesMaster1.xml.rels"] = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="../theme/theme2.xml"/>
</Relationships>
""".encode()

    rels = etree.fromstring(files["ppt/_rels/presentation.xml.rels"])
    ids = [int(rel.get("Id")[3:]) for rel in rels if rel.get("Id", "").startswith("rId")]
    next_id = max(ids) + 1
    etree.SubElement(rels, f"{{{nsr}}}Relationship", {
        "Id": f"rId{next_id}",
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesMaster",
        "Target": "notesMasters/notesMaster1.xml",
    })
    files["ppt/_rels/presentation.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)

    pres = etree.fromstring(files["ppt/presentation.xml"])
    notes_list = etree.Element(f"{{{nsp}}}notesMasterIdLst")
    etree.SubElement(notes_list, f"{{{nsp}}}notesMasterId").set(f"{{{nsr_attr}}}id", f"rId{next_id}")
    pres.find(f"{{{nsp}}}sldMasterIdLst").addnext(notes_list)
    files["ppt/presentation.xml"] = etree.tostring(pres, xml_declaration=True, encoding="UTF-8", standalone=True)

    content_types = etree.fromstring(files["[Content_Types].xml"])

    def override(part, ctype):
        etree.SubElement(content_types, f"{{{nsct}}}Override", {"PartName": part, "ContentType": ctype})

    override("/ppt/theme/theme2.xml", "application/vnd.openxmlformats-officedocument.theme+xml")
    override("/ppt/notesMasters/notesMaster1.xml", "application/vnd.openxmlformats-officedocument.presentationml.notesMaster+xml")

    for index, text in enumerate(notes, 1):
        files[f"ppt/notesSlides/notesSlide{index}.xml"] = notes_xml(text)
        files[f"ppt/notesSlides/_rels/notesSlide{index}.xml.rels"] = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="../slides/slide{index}.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesMaster" Target="../notesMasters/notesMaster1.xml"/>
</Relationships>
""".encode()
        override(
            f"/ppt/notesSlides/notesSlide{index}.xml",
            "application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml",
        )
        slide_rels = etree.fromstring(files[f"ppt/slides/_rels/slide{index}.xml.rels"])
        slide_ids = [int(rel.get("Id")[3:]) for rel in slide_rels if rel.get("Id", "").startswith("rId")]
        etree.SubElement(slide_rels, f"{{{nsr}}}Relationship", {
            "Id": f"rId{max(slide_ids) + 1}",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide",
            "Target": f"../notesSlides/notesSlide{index}.xml",
        })
        files[f"ppt/slides/_rels/slide{index}.xml.rels"] = etree.tostring(
            slide_rels, xml_declaration=True, encoding="UTF-8", standalone=True
        )

    files["[Content_Types].xml"] = etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)


if __name__ == "__main__":
    shots = prepare_images()
    path = build(shots)
    md = write_narracion()
    rewrite_notes_for_keynote(path)
    pdf = export_pdf(path)
    print(path)
    print(pdf)
    print(md)
